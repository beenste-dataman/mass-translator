import os
import argparse
import sys
from tqdm import tqdm
import xlrd
import xlwt
# Required libraries
import docx
import openpyxl
import csv
import json
import torch
from bs4 import BeautifulSoup
from transformers import MarianMTModel, MarianTokenizer
from concurrent.futures import ThreadPoolExecutor
from odf.opendocument import OpenDocumentText
from odf.text import P, Span
from odf import teletype

art = '''
                                 __                              .__          __                
  _____ _____    ______ ______ _/  |_____________    ____   _____|  | _____ _/  |_  ___________ 
 /     \\__  \  /  ___//  ___/ \   __\_  __ \__  \  /    \ /  ___/  | \__  \\   __\/  _ \_  __ \
|  Y Y  \/ __ \_\___ \ \___ \   |  |  |  | \// __ \|   |  \\___ \|  |__/ __ \|  | (  <_> )  | \/
|__|_|  (____  /____  >____  >  |__|  |__|  (____  /___|  /____  >____(____  /__|  \____/|__|   
      \/     \/     \/     \/                    \/     \/     \/          \/                   
'''

print(art)



def parse_args():
    parser = argparse.ArgumentParser(description="Translate files in a directory and its subdirectories.")
    parser.add_argument("input_dir", help="The path to the input directory.")
    parser.add_argument("output_dir", help="The path to the output directory.")
    parser.add_argument("--src_language", default="en", help="The source language (default: 'en').")
    parser.add_argument("--tgt_language", default="es", help="The target language (default: 'es').")
    return parser.parse_args()

# Initialize the model and tokenizer
args = parse_args()
model_name = f"Helsinki-NLP/opus-mt-{args.src_language}-{args.tgt_language}"
tokenizer = MarianTokenizer.from_pretrained(model_name)
model = MarianMTModel.from_pretrained(model_name)


#gpu
device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
model.to(device)



# Translation function
def translate_text(text):
    # If the text is empty or contains only whitespace, return it as-is
    if not text.strip():
        return text

    # Split the input text into chunks of approximately 100 words
    text_chunks = text.split(' ')
    chunk_size = 100
    chunks = [' '.join(text_chunks[i:i + chunk_size]) for i in range(0, len(text_chunks), chunk_size)]

    translated_chunks = []
    for chunk in chunks:
        print(f"Translating: {chunk[:50]}...")  # progress verbosity

        # Tokenize and translate the chunk
        tokenized_chunk = tokenizer(chunk, return_tensors="pt")
        input_ids = tokenized_chunk["input_ids"].to(device)  # Move input tensor to the device

        translated_tokens = model.generate(input_ids)
        translated_chunk = tokenizer.decode(translated_tokens[0], skip_special_tokens=True)

        print(f"Translated: {translated_chunk[:50]}...")  # further verbosity
        translated_chunks.append(translated_chunk)

    # Combine the translated chunks
    translated_text = ' '.join(translated_chunks)
    return translated_text



  
# File processing function
def process_files(input_dir, output_dir):
    files_to_process = []
    for root, _, files in os.walk(input_dir):
        for file in files:
            file_path = os.path.join(root, file)
            files_to_process.append(file_path)

    progress_bar = tqdm(total=len(files_to_process), desc="Processing files", unit="file")

    def process_file_with_progress(file_path):
        try:
            translate_file(file_path)
        except Exception as e:
            progress_bar.write(f"Error processing {file_path}: {e}")
        finally:
            progress_bar.update()
# set threadpoolexecuter max_workers=
    with ThreadPoolExecutor(max_workers=10) as executor:
        executor.map(process_file_with_progress, files_to_process)

    progress_bar.close()
    
    
    
    
# Get filey    
    

def translate_file(file_path):
    file_ext = os.path.splitext(file_path)[1]
    if file_ext == ".docx":
        translate_docx(file_path)
    elif file_ext == ".xlsx":
        translate_xlsx(file_path)
    elif file_ext == ".xls":
        translate_xls(file_path)
    elif file_ext == ".csv":
        translate_csv(file_path)
    elif file_ext == ".json":
        translate_json(file_path)
    elif file_ext == ".html":
        translate_html(file_path)
    elif file_ext == ".odt":
        translate_odt(file_path)
    elif file_ext == ".txt":
        translate_txt(file_path)
    
    else:
        print(f"Unsupported file type: {file_ext}")

        
# file handlers

def translate_dict(obj):
    if isinstance(obj, dict):
        return {key: translate_dict(value) for key, value in obj.items()}
    elif isinstance(obj, list):
        return [translate_dict(item) for item in obj]
    elif isinstance(obj, str):
        return translate_text(obj)
    else:
        return obj

   
# Define translation functions for each file type
def translate_docx(file_path):
    doc = docx.Document(file_path)
    for paragraph in doc.paragraphs:
        translated_text = translate_text(paragraph.text)
        paragraph.text = translated_text

    output_path = os.path.join(output_dir, os.path.relpath(file_path, input_dir))
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)

def translate_xlsx(file_path):
    workbook = openpyxl.load_workbook(file_path)
    for sheet in workbook.worksheets:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    translated_text = translate_text(cell.value)
                    cell.value = translated_text

    output_path = os.path.join(output_dir, os.path.relpath(file_path, input_dir))
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    workbook.save(output_path)

def translate_csv(file_path):
    with open(file_path, "r", encoding="utf-8") as file:
        reader = csv.reader(file)
        rows = list(reader)

    translated_rows = []
    for row in rows:
        translated_row = [translate_text(cell) for cell in row]
        translated_rows.append(translated_row)

    output_path = os.path.join(output_dir, os.path.relpath(file_path, input_dir))
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerows(translated_rows)

def translate_json(file_path):
    with open(file_path, "r", encoding="utf-8") as file:
        data = json.load(file)

    translated_data = translate_dict(data)

    output_path = os.path.join(output_dir, os.path.relpath(file_path, input_dir))
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as file:
        json.dump(translated_data, file, ensure_ascii=False, indent=2)

def translate_html(file_path):
    with open(file_path, "r", encoding="utf-8") as file:
        soup = BeautifulSoup(file.read(), "html.parser")

    translate_soup(soup)

    output_path = os.path.join(output_dir, os.path.relpath(file_path, input_dir))
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as file:
        file.write(str(soup))


def translate_xls(file_path):
    workbook = xlrd.open_workbook(file_path)
    translated_workbook = xlwt.Workbook()

    for sheet_idx in range(workbook.nsheets):
        sheet = workbook.sheet_by_index(sheet_idx)
        translated_sheet = translated_workbook.add_sheet(sheet.name)

        for row_idx in range(sheet.nrows):
            for col_idx in range(sheet.ncols):
                cell_value = sheet.cell_value(row_idx, col_idx)

                if isinstance(cell_value, str):
                    translated_text = translate_text(cell_value)
                    translated_sheet.write(row_idx, col_idx, translated_text)
                else:
                    translated_sheet.write(row_idx, col_idx, cell_value)

    output_path = os.path.join(output_dir, os.path.relpath(file_path, input_dir))
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    translated_workbook.save(output_path)

    
def translate_odt(file_path):
    document = OpenDocumentText(file_path)
    translated_document = OpenDocumentText()

    for element in document.getElementsByType(P):
        paragraph = P()
        for child in element.childNodes:
            if child.qname[1] == "span":
                translated_text = translate_text(teletype.extractText(child))
                span = Span()
                span.addText(translated_text)
                paragraph.addElement(span)
            elif isinstance(child, str):
                translated_text = translate_text(child)
                paragraph.addText(translated_text)
        translated_document.text.addElement(paragraph)

    output_path = os.path.join(output_dir, os.path.relpath(file_path, input_dir))
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    translated_document.save(output_path)   
    
    
def translate_txt(file_path):
    with open(file_path, "r", encoding="utf-8") as file:
        content = file.read()

    translated_content = translate_text(content)

    output_path = os.path.join(output_dir, os.path.relpath(file_path, input_dir))
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as output_file:
        output_file.write(translated_content)

        
        
        
        
        
        
# Rest of the code remains the same

if __name__ == "__main__":
    input_dir = args.input_dir
    output_dir = args.output_dir

    if not os.path.exists(input_dir):
        print(f"Input directory '{input_dir}' does not exist.")
        sys.exit(1)

    os.makedirs(output_dir, exist_ok=True)
    process_files(input_dir, output_dir)
